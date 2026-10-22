#!/usr/bin/env python3
"""Puts the report in the same order as the notebook: ... Results, Green AI (now Section 10), then
Scoring a new company (now Section 11). Renumbers the two sections, their tables (Green AI 20-22,
scoring 23-24) and figures (Green AI 10, scoring 11), and every cross-reference to them.

    python reorder_sections.py IN.docx OUT.docx
"""
import re, sys
import docx
from docx.oxml.ns import qn

src, dst = sys.argv[1:3]
d = docx.Document(src)
body = d.element.body
txt = lambda el: "".join(t.text or "" for t in el.iter(qn("w:t")))
els = list(body.iterchildren())
i_score = next(i for i, e in enumerate(els) if txt(e).startswith("10. Scoring a New Company"))
i_green = next(i for i, e in enumerate(els) if txt(e).startswith("11. Green AI"))
i_concl = next(i for i, e in enumerate(els) if txt(e).startswith("12. Conclusion"))
assert i_score < i_green < i_concl
for e in els[i_score:i_green]:              # the scoring section goes after Green AI
    els[i_concl].addprevious(e)

MAP = {"10. Scoring a New Company": "11. Scoring a New Company",
       "11. Green AI": "10. Green AI",
       "Section 10": "Section 11", "Section 11": "Section 10",
       "Table 20": "Table 23", "Table 21": "Table 24",
       "Table 22": "Table 20", "Table 23": "Table 21", "Table 24": "Table 22",
       "Figure 10": "Figure 11", "Figure 11": "Figure 10"}
pat = re.compile("|".join(re.escape(k) for k in sorted(MAP, key=len, reverse=True)))
for t in body.iter(qn("w:t")):
    if t.text:
        t.text = pat.sub(lambda m: MAP[m.group(0)], t.text)

# the date line in the information table
for row in d.tables[0].rows:
    if row.cells[0].text.strip() == "Report date":
        p = row.cells[1].paragraphs[0]
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = "22 October 2026 (first issued 3 October; Green AI study added as Section 10)"
d.save(dst)
print("saved", dst)
